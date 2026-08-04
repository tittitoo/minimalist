"""
Measure the TRUE line count of every wrapped Description cell in a generated proposal,
by reading the real rendered geometry out of the PDF — then fit _SP_MDW_PX to it.

Why this exists
---------------
_SP_MDW_PX converts an Excel column_width into the text width Excel's PDF renderer
actually gives you.  It was historically calibrated by eyeballing a handful of rows in
`pdftotext -layout` output, which is unreliable: pdftotext reports a hyphen-wrapped word
("electro-" / "polished") or trailing punctuation as though text were missing, so cases
that rendered perfectly looked like clipping bugs.  Several rounds of calibration chased
those false positives and made the constant worse.

This script measures instead of guessing.  It pulls real glyph x-coordinates from the PDF
(via PyMuPDF), reconstructs which rendered lines belong to which source cell, and reports
the true line count per row.  Feeding that to a sweep over avail_pt gives the range of
values that reproduce reality exactly.

Usage
-----
    uv pip install pymupdf          # not a runtime dependency; analysis only
    python tools/extract_wrap_ground_truth.py <proposal.xlsx> <proposal.pdf> <col_width>

e.g.
    python tools/extract_wrap_ground_truth.py "Commercial J12632 ....xlsx" \
                                              "Commercial J12632 ....pdf" 55

Reading the output
------------------
"perfect avail range" is the set of avail_pt values with zero mismatches; the implied MDW
range is printed alongside.  Run it for BOTH the Commercial (col=55) and Technical
(col=68) outputs and intersect the two windows — agreement across two column widths is
what makes a calibration trustworthy.

IMPORTANT CAVEAT: rows that are actually CLIPPED cannot be matched (their rendered text
does not equal the source cell), so they are silently excluded from the ground-truth set.
A perfect fit here therefore does NOT by itself prove clipped rows are fixed.  Add each
known clip case as an explicit extra constraint when picking the final value — see the
"(REMOVED)" case referenced in the _SP_MDW_PX comment in functions.py.
"""

import os
import re
import sys

import fitz  # PyMuPDF
import openpyxl
from reportlab.pdfbase.pdfmetrics import stringWidth as sw

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

FONT, PT = "Helvetica", 12  # metrically equivalent to Excel's Arial

_norm = lambda s: re.sub(r"\s+", " ", s).strip()


def rendered_lines(pdf_path, x_max, x_min=80):
    """Rendered text lines inside the Description column, in reading order."""
    doc = fitz.open(pdf_path)
    out = []
    for page in doc:
        items = []
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                text = "".join(s["text"] for s in line["spans"])
                if not text.strip():
                    continue
                x0 = min(s["bbox"][0] for s in line["spans"])
                x1 = max(s["bbox"][2] for s in line["spans"])
                y0 = min(s["bbox"][1] for s in line["spans"])
                # Column C only, and drop the page-number footer.
                if x0 >= x_min and x1 <= x_max and not re.match(r"^\s*Page \d+ of \d+\s*$", text):
                    items.append((y0, text))
        items.sort(key=lambda t: t[0])
        out.extend(t for _, t in items)
    return out


def source_rows(xlsx_path, first_row=18):
    ws = openpyxl.load_workbook(xlsx_path, data_only=True)["Proposal"]
    return [
        (r, str(ws.cell(row=r, column=3).value))
        for r in range(first_row, ws.max_row + 1)
        if ws.cell(row=r, column=3).value and str(ws.cell(row=r, column=3).value).strip()
    ]


def match_line_counts(rows, lines):
    """Map each source row to the number of rendered lines it occupies.

    Rows whose rendered text does not reassemble exactly (i.e. clipped rows) are
    skipped rather than guessed at — see the caveat in the module docstring.
    """
    truth, i = {}, 0
    for row_num, text in rows:
        target = _norm(text)
        if i >= len(lines):
            break
        # Resync if we've drifted out of step with the rendered stream.
        if not target.startswith(_norm(lines[i])[:25]):
            j = i
            while j < min(i + 40, len(lines)) and not target.startswith(_norm(lines[j])[:25]):
                j += 1
            if j >= min(i + 40, len(lines)):
                continue
            i = j
        acc, n = "", 0
        while i < len(lines) and n < 15:
            piece = lines[i]
            if not acc:
                acc = piece
            elif acc.rstrip().endswith("-") and not acc.endswith(" "):
                acc = acc[:-1] + piece      # hyphen-wrapped word rejoins with no space
            else:
                acc = acc + " " + piece
            n += 1
            i += 1
            if _norm(acc) == target:
                truth[row_num] = n
                break
            if not target.startswith(_norm(acc)[: max(1, len(_norm(acc)) - 3)]):
                break
    return truth


def wrap_lines(text, avail_pt):
    """Mirror of functions._sp_wrap_lines — keep the two in sync when either changes."""
    total = 0
    for segment in text.split("\n"):
        if not segment.strip():
            total += 1
            continue
        body = segment.lstrip()
        lead_w = sw(segment[: len(segment) - len(body)], FONT, PT)
        sp_w = sw(" ", FONT, PT)
        seg_lines, cur = 1, None
        for word in body.split():
            ww = sw(word, FONT, PT)
            if cur is None:
                cur = lead_w + ww
            elif cur + sp_w + ww > avail_pt:
                seg_lines += 1
                cur = ww
            else:
                cur += sp_w + ww
        total += seg_lines
    return max(1, total)


def main():
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    xlsx_path, pdf_path, col_width = sys.argv[1], sys.argv[2], float(sys.argv[3])

    # Description column runs from ~x=80 to just left of the Qty column; widen the
    # window for the wider Technical layout.
    x_max = 350 if col_width < 60 else 420

    rows = source_rows(xlsx_path)
    truth = match_line_counts(rows, rendered_lines(pdf_path, x_max))
    row_text = dict(rows)
    print(f"matched {len(truth)} of {len(rows)} rows against real rendered geometry")
    if not truth:
        print("no rows matched — check the sheet name / column range")
        return

    good = [
        a for a in range(250, 460)
        if all(wrap_lines(row_text[r], float(a)) == n for r, n in truth.items())
    ]
    if good:
        lo, hi = min(good), max(good)
        print(f"perfect avail range: {lo}-{hi}pt")
        print(f"  => implied _SP_MDW_PX {(lo/0.75-1)/col_width:.3f} - {(hi/0.75-1)/col_width:.3f}")
        print("  intersect this with the other column width's range before choosing.")
    else:
        print("no single avail_pt reproduces every row — investigate the mismatches below")

    from functions import _SP_MDW_PX
    cur = (col_width * _SP_MDW_PX + 1) * 0.75
    bad = [(r, n, wrap_lines(row_text[r], cur)) for r, n in truth.items()
           if wrap_lines(row_text[r], cur) != n]
    print(f"\ncurrent _SP_MDW_PX={_SP_MDW_PX} (avail={cur:.1f}pt): {len(bad)} mismatches")
    for r, true_n, pred_n in bad[:20]:
        print(f"  row {r}: true={true_n} predicted={pred_n} :: {row_text[r][:70]!r}")


if __name__ == "__main__":
    main()
